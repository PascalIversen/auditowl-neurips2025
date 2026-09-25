#!/usr/bin/env python3
"""Section 4.4 - the GPT reviewer's verdicts, over every finding it judged.

The GPT reviewer was given the same instructions as the human experts and run
over the same 40 papers as the human evaluation study (a wider finding pool:
not every GPT-judged finding was also human-rated). gpt_ratings.xlsx and
gpt_ratings_round2.xlsx are disjoint batches (87 + 75 sheets' worth of rows,
0 overlap), so no dedup across them is needed.

Run: python studies/expert_review/gpt_reviewer_verdicts.py
"""
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent

WORKBOOKS = ("gpt_ratings.xlsx", "gpt_ratings_round2.xlsx")
COL_NUM, COL_ISSUE, COL_CORRECTNESS, COL_RELEVANCE = 1, 2, 10, 12
FIRST_DATA_ROW = 5


def main():
    correctness, relevance = Counter(), Counter()
    n = 0

    for name in WORKBOOKS:
        wb = openpyxl.load_workbook(HERE / name, data_only=True)
        for ws in wb.worksheets:
            if ws.title == "Instructions":
                continue
            for row in range(FIRST_DATA_ROW, ws.max_row + 1):
                num = ws.cell(row=row, column=COL_NUM).value
                if num is None or not str(num).strip().replace(".", "").isdigit():
                    continue
                n += 1
                for col, tally in ((COL_CORRECTNESS, correctness), (COL_RELEVANCE, relevance)):
                    v = ws.cell(row=row, column=col).value
                    if v:
                        tally[str(v).strip()] += 1

    print(f"GPT reviewer verdicts on all {n} findings it judged\n")
    for label, tally in (("correctness", correctness), ("relevance", relevance)):
        print(f"  {label}")
        for verdict, n_v in tally.most_common():
            print(f"    {verdict:<28}{n_v:>4}")


if __name__ == "__main__":
    main()
