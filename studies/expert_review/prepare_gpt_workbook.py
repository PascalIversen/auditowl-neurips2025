#!/usr/bin/env python3
"""Make the blinded GPT copy of the adapted evaluation workbook.

- clears every human verdict/input cell on the paper sheets (rows >= 5, cols 9-14)
- sets each paper sheet's reviewer field to "GPT"
- Instructions sheet: replaces human names in the WHO DOES WHAT table with "Done"
  (header renamed to "Status") so the GPT run carries no names and no verdicts

Input:  ~/Downloads/AuditOwl Evaluation overview.xlsx  (the humans' live workbook)
Output: gpt_reviewer/gpt_workbook.xlsx
"""
import os

from openpyxl import load_workbook

SRC = os.path.expanduser("~/Downloads/AuditOwl Evaluation overview.xlsx")
HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "gpt_workbook.xlsx")


def main():
    wb = load_workbook(SRC)
    ins = wb["Instructions"]
    # WHO DOES WHAT table: header row 2 (C=Paper, D=Title, E=names)
    hdr = ins.cell(row=2, column=5)
    if hdr.value:
        hdr.value = "Status"
    for r in range(3, ins.max_row + 1):
        c = ins.cell(row=r, column=5)
        if c.value not in (None, ""):
            c.value = "Done"
    for name in wb.sheetnames:
        if name == "Instructions":
            continue
        ws = wb[name]
        ws.cell(row=2, column=1).value = "Reviewer name: GPT"
        ws.cell(row=2, column=2).value = None
        for r in range(5, ws.max_row + 1):
            if ws.cell(row=r, column=1).value in (None, ""):
                continue
            for c in range(9, 15):
                cell = ws.cell(row=r, column=c)
                cell.value = None
                cell.hyperlink = None
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
