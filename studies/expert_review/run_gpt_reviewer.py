#!/usr/bin/env python3
"""GPT reviewer arm: the identical task the humans do, run via the OpenAI Codex CLI.

For each paper sheet in gpt_workbook.xlsx:
  - working dir = ../reviewer_input/<paper_dir>/ (paper.pdf, paper_text.txt, code/ --
    exactly the human hand-out; no verdict files exist there, so the run is blinded)
  - prompt = the verbatim instructions from the workbook's Instructions sheet
    + the paper's issue cards (read from the sheet, so GPT sees what humans see)
    + a JSON output contract
  - the model writes gpt_verdicts.json in its working dir; the driver validates it.

Commands:
  run  [--papers 205 570 ...] [--model gpt-5.5] [--effort high] [--jobs 3]
  fill                       # write all gpt_verdicts.json into gpt_workbook.xlsx

Per-paper wall time is recorded by the driver and written next to the reviewer name
(the model's own "minutes" are not trusted and not written into the Minutes column).
Resume via done.txt; logs in logs/.
"""
import argparse
import glob
import json
import os
import shutil
import subprocess
import time
from concurrent.futures import ThreadPoolExecutor, as_completed

from openpyxl import load_workbook

def _find_root(start):
    """Repo root = the ancestor holding audits/. Found by content rather than a
    fixed number of dirname() hops, so this survives the file moving depth."""
    cur = start
    while True:
        if os.path.isdir(os.path.join(cur, "audits")):
            return cur
        parent = os.path.dirname(cur)
        if parent == cur:
            raise SystemExit(f"no `audits/` directory found above {start}")
        cur = parent


HERE = os.path.dirname(os.path.abspath(__file__))
EXP = os.path.dirname(HERE)
ROOT = _find_root(HERE)
RINPUT = os.path.join(EXP, "reviewer_input")
WB = os.path.join(HERE, "gpt_workbook.xlsx")
LOGS = os.path.join(HERE, "logs")
DONE = os.path.join(HERE, "done.txt")

CORRECTNESS = ["True", "Partially true", "False", "Unverifiable"]
RELEVANCE = ["Relevant", "Minor", "Trivial", "Unverifiable", "N/A (correctness False)"]

# Papers whose audited artifact is NOT code/<owner>__<repo>/. 3463's findings were
# written against the NeurIPS supplemental ZIP (audit.md:14-24); code/ holds the
# authors' later public GitHub release, a different revision on which the cited line
# numbers do not resolve. See reviewer_input/<paper>/ARTIFACT_NOTE.md.
CODE_DIR = {"3463": "code_supplemental"}

CONTRACT = """
---

OUTPUT CONTRACT (driver-facing; this replaces filling the Excel by hand)

You are the reviewer. Work through the issues exactly as instructed above: for each
issue open the cited file and line in ./__CODE_DIR__, compare with the paper (paper.pdf;
paper_text.txt is the extracted text), and judge both axes. Only use the provided
materials; no internet.

When done, write a file `gpt_verdicts.json` in the current directory:

[
  {"issue": <number from the card>,
   "correctness": one of ["True", "Partially true", "False", "Unverifiable"],
   "reason_correctness": "<short reason; REQUIRED if correctness is False or Unverifiable, else may be empty>",
   "relevance": one of ["Relevant", "Minor", "Trivial", "Unverifiable", "N/A (correctness False)"],
   "reason_relevance": "<short reason; REQUIRED if relevance is Trivial or Unverifiable, else may be empty>",
   "notes": "<optional>"},
  ... one object per issue, in order ...
]

Use "N/A (correctness False)" for relevance if and only if correctness is "False".
The file must be valid JSON, nothing else in it.
"""


def sheet_instructions(wb):
    ws = wb["Instructions"]
    lines = []
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        lines.append("" if v is None else str(v))
    while lines and not lines[-1]:
        lines.pop()
    return "\n".join(lines)


def sheet_issues(ws):
    hdr = [ws.cell(row=4, column=c).value for c in range(1, 9)]
    items = []
    for r in range(5, ws.max_row + 1):
        num = ws.cell(row=r, column=1).value
        if num in (None, ""):
            continue
        vals = [ws.cell(row=r, column=c).value for c in range(1, 9)]
        card = [f"ISSUE {int(float(vals[0]))}"]
        for h, v in zip(hdr[1:], vals[1:]):
            card.append(f"{h}: {v if v is not None else ''}")
        items.append("\n".join(card))
    return items


def paper_dir(pid):
    hits = [d for d in glob.glob(os.path.join(RINPUT, f"{pid}_*")) if os.path.isdir(d)]
    assert len(hits) == 1, f"no unique reviewer_input dir for {pid}: {hits}"
    return hits[0]


def ensure_paper_text(pid, pdir):
    dst = os.path.join(pdir, "paper_text.txt")
    if not os.path.exists(dst):
        src = glob.glob(os.path.join(ROOT, "audits", f"{pid}", "paper_text.txt"))
        if src:
            shutil.copy(src[0], dst)


def build_prompt(instructions, cards, code_dir="code"):
    head = ("You are one of the reviewers in this study. Below are the study "
            "instructions given to every reviewer, followed by the issue cards for the "
            "paper in your current directory.\n\n")
    contract = CONTRACT.replace("__CODE_DIR__", code_dir)
    if code_dir != "code":
        contract += ("\nThis paper ships more than one code tree. Read ARTIFACT_NOTE.md in "
                     f"the current directory first: the issue cards' line numbers are "
                     f"relative to ./{code_dir}, which is the audited artifact.\n")
    return (head + instructions + "\n\n---\n\nISSUE CARDS\n\n"
            + "\n\n".join(cards) + "\n" + contract)


def validate(path, n_expected):
    data = json.load(open(path))
    assert isinstance(data, list) and len(data) == n_expected, \
        f"expected {n_expected} verdicts, got {len(data) if isinstance(data, list) else type(data)}"
    for d in data:
        assert d.get("correctness") in CORRECTNESS, d
        assert d.get("relevance") in RELEVANCE, d
        if d["correctness"] in ("False", "Unverifiable"):
            assert str(d.get("reason_correctness") or "").strip(), f"missing reason_correctness: {d}"
        if d["relevance"] in ("Trivial", "Unverifiable"):
            assert str(d.get("reason_relevance") or "").strip(), f"missing reason_relevance: {d}"
        assert (d["relevance"] == "N/A (correctness False)") == (d["correctness"] == "False"), d
    return data


def run_paper(pid, instructions, cards, model, effort, timeout):
    pdir = paper_dir(pid)
    ensure_paper_text(pid, pdir)
    out = os.path.join(pdir, "gpt_verdicts.json")
    if os.path.exists(out):
        os.remove(out)
    prompt = build_prompt(instructions, cards, CODE_DIR.get(pid, "code"))
    log = os.path.join(LOGS, f"{pid}.jsonl")
    cmd = ["codex", "exec", "-m", model, "-C", pdir,
           "--sandbox", "workspace-write", "--skip-git-repo-check",
           "--color", "never", "--json",
           "-c", f'model_reasoning_effort="{effort}"',
           "-o", os.path.join(LOGS, f"{pid}.last.md"), "-"]
    t0 = time.time()
    with open(log, "w") as lf:
        r = subprocess.run(cmd, input=prompt, stdout=lf, stderr=subprocess.STDOUT,
                           text=True, timeout=timeout)
    minutes = (time.time() - t0) / 60
    validate(out, len(cards))
    with open(os.path.join(HERE, f"meta_{pid}.json"), "w") as fh:
        json.dump({"paper": pid, "model": model, "effort": effort,
                   "code_dir": CODE_DIR.get(pid, "code"),
                   "wall_minutes": round(minutes, 1), "returncode": r.returncode}, fh)
    return pid, minutes


def apply_paths(args):
    global WB, RINPUT
    if getattr(args, "wb", None):
        WB = os.path.abspath(args.wb)
    if getattr(args, "rinput", None):
        RINPUT = os.path.abspath(args.rinput)
    missing = [p for p in (WB, RINPUT) if not os.path.exists(p)]
    if missing:
        raise SystemExit(
            "missing input(s), not part of this release: " + ", ".join(missing) +
            " -- this arm's outputs (gpt_ratings*.xlsx) are released and "
            "gpt_reviewer_verdicts.py recomputes §4.4's numbers from them; "
            "only re-executing this driver requires the workbook.")


def cmd_run(args):
    apply_paths(args)
    os.makedirs(LOGS, exist_ok=True)
    wb = load_workbook(WB)
    instructions = sheet_instructions(wb)
    done = set()
    if os.path.exists(DONE):
        done = set(open(DONE).read().split())
    papers = args.papers or [s for s in wb.sheetnames if s != "Instructions"]
    todo = [p for p in papers if p not in done]
    print(f"{len(todo)} papers to run: {todo}")
    tasks = {p: sheet_issues(wb[p]) for p in todo}
    with ThreadPoolExecutor(max_workers=args.jobs) as ex:
        futs = {ex.submit(run_paper, p, instructions, tasks[p],
                          args.model, args.effort, args.timeout): p for p in todo}
        for f in as_completed(futs):
            p = futs[f]
            try:
                _, minutes = f.result()
                with open(DONE, "a") as fh:
                    fh.write(p + "\n")
                print(f"  done {p} ({minutes:.0f} min)")
            except Exception as e:
                print(f"  FAILED {p}: {e}")


def cmd_fill(args):
    apply_paths(args)
    wb = load_workbook(WB)
    filled = 0
    for pid in [s for s in wb.sheetnames if s != "Instructions"]:
        out = os.path.join(paper_dir(pid), "gpt_verdicts.json")
        if not os.path.exists(out):
            continue
        ws = wb[pid]
        cards = sheet_issues(ws)
        data = validate(out, len(cards))
        meta_p = os.path.join(HERE, f"meta_{pid}.json")
        if os.path.exists(meta_p):
            meta = json.load(open(meta_p))
            ws.cell(row=2, column=4, value=f"GPT wall time: {meta['wall_minutes']} min "
                    f"({meta['model']}, {meta['effort']})")
        by_issue = {d["issue"]: d for d in data}
        for r in range(5, ws.max_row + 1):
            num = ws.cell(row=r, column=1).value
            if num in (None, ""):
                continue
            d = by_issue[int(float(num))]
            ws.cell(row=r, column=10, value=d["correctness"])
            ws.cell(row=r, column=11, value=d.get("reason_correctness") or None)
            ws.cell(row=r, column=12, value=d["relevance"])
            ws.cell(row=r, column=13, value=d.get("reason_relevance") or None)
            ws.cell(row=r, column=14, value=d.get("notes") or None)
        filled += 1
    wb.save(WB)
    print(f"filled {filled} paper sheets into {WB}")


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--papers", nargs="*")
    r.add_argument("--model", default="gpt-5.5")
    r.add_argument("--effort", default="high")
    r.add_argument("--jobs", type=int, default=3)
    r.add_argument("--timeout", type=int, default=5400)
    r.add_argument("--wb", help="workbook to read cards from (default batch-1 gpt_workbook.xlsx)")
    r.add_argument("--rinput", help="reviewer_input dir (default batch-1)")
    r.set_defaults(fn=cmd_run)
    f = sub.add_parser("fill")
    f.add_argument("--wb")
    f.add_argument("--rinput")
    f.set_defaults(fn=cmd_fill)
    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
