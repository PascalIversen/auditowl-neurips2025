#!/usr/bin/env python3
"""Anonymised aggregate of the author self-evaluation (Section 4.2 / S1.3).

We asked the authors of every sampled paper to judge the discrepancies raised
against their own work. The raw responses are not published: they are written
in the authors' own words and attributable to a specific paper, and the Ethics
Statement commits to reporting them in aggregate only.

This emits the aggregate the reported numbers are computed from: one row per
judged discrepancy carrying only the finding's own category and the two
verdicts. No paper identifier, no issue id, no free text, and rows are sorted
by content so their order carries no grouping. From this the Section S1.3
counts are recomputable; individual responses are not recoverable.

Run: python studies/expert_review/build_author_response_summary.py
"""
import csv
import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
WB = HERE / "author_responses.xlsx"
OUT_CSV = HERE / "author_response_verdicts.csv"
OUT_JSON = HERE / "author_response_summary.json"

# the authors' free-text variants, folded onto the scale they were offered
RELEVANCE_SYNONYM = {
    "Minor": "Minor point",
    "Medium": "Minor point",
    "Irrelevant / Minor point": "Irrelevant",
    "N/A (correctness False)": "",
}
COL = {"issue": 1, "category": 3, "correctness": 10, "relevance": 12}


def norm(v):
    return str(v).strip() if v is not None and str(v).strip() else ""


def main():
    if not WB.exists():
        raise SystemExit(
            f"{WB.name} not found. The response workbook is not part of the "
            "release; this script regenerates the aggregate from the authors' "
            "local copy."
        )
    wb = load_workbook(WB, data_only=True)

    rows, papers_with_verdicts = [], 0
    for sheet in wb.sheetnames:
        if sheet in ("Instructions", "Overview"):
            continue
        ws = wb[sheet]
        header = next((r for r in range(1, 12)
                       if norm(ws.cell(row=r, column=1).value) == "#"), None)
        if header is None:
            continue
        judged_here = 0
        for r in range(header + 1, ws.max_row + 1):
            num = norm(ws.cell(row=r, column=COL["issue"]).value)
            if not num or not num.replace(".", "").isdigit():
                break                      # findings table ends here
            correctness = norm(ws.cell(row=r, column=COL["correctness"]).value)
            if not correctness:
                continue                   # finding not explicitly addressed
            relevance = norm(ws.cell(row=r, column=COL["relevance"]).value)
            rows.append({
                "category": norm(ws.cell(row=r, column=COL["category"]).value),
                "author_correctness": correctness,
                "author_relevance": RELEVANCE_SYNONYM.get(relevance, relevance),
            })
            judged_here += 1
        papers_with_verdicts += judged_here > 0

    # sorting by content removes any residual per-paper ordering
    rows.sort(key=lambda r: (r["category"], r["author_correctness"], r["author_relevance"]))

    with OUT_CSV.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)

    correctness = Counter(r["author_correctness"] for r in rows)
    summary = {
        "papers_replying_with_per_finding_verdicts": papers_with_verdicts,
        "findings_judged": len(rows),
        "correctness": dict(correctness),
        "correct_or_partially_correct": correctness["Correct"] + correctness["Partially correct"],
        "relevance": dict(Counter(r["author_relevance"] for r in rows if r["author_relevance"])),
        "by_category": {c: dict(Counter(r["author_correctness"] for r in rows if r["category"] == c))
                        for c in sorted({r["category"] for r in rows})},
    }
    OUT_JSON.write_text(json.dumps(summary, indent=2) + "\n")

    print(f"wrote {OUT_CSV.name} ({len(rows)} judged discrepancies) and {OUT_JSON.name}")
    print(f"  {papers_with_verdicts} papers replied with per-finding verdicts")
    print(f"  {summary['correct_or_partially_correct']} of {len(rows)} correct or partially correct "
          f"({correctness['Correct']} fully, {correctness['Partially correct']} partially); "
          f"{correctness['Incorrect']} rejected")


if __name__ == "__main__":
    main()
