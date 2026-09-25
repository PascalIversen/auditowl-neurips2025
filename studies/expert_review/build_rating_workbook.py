#!/usr/bin/env python3
"""Build the reviewer Excel workbook: Instructions sheet + one sheet per covered paper.

Sheet 1 explains the study and the task. Each paper sheet lists the paper's in-scope
issues with dropdown verdict cells (Correctness, Relevance), a minutes field, and free-text
reason/notes, plus a reviewer-name field. Output: _selection/review_workbook.xlsx
"""
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SEL = os.path.join(HERE, "_selection")
OUT = os.path.join(SEL, "review_workbook.xlsx")

IN_SCOPE = {
    ("high", "missing"), ("high", "difference"), ("high", "bug"), ("high", "methodology"),
    ("medium", "difference"), ("medium", "bug"), ("medium", "methodology"),
}
CATEGORY_LABEL = {"difference": "mismatch"}

CORRECTNESS = ["True", "Partially true", "False", "Unverifiable"]
RELEVANCE = ["Relevant", "Minor", "Trivial", "Unverifiable", "N/A (correctness False)"]

INSTRUCTIONS = [
    ("AuditOwl human validation: instructions", True),
    ("", False),
    ("Thanks for helping! Since reviewers rarely have time to review code alongside papers, "
     "we built a tool that reads a computational paper and its released code and flags "
     "\"discrepancies\": places where the code is broken or does not support what the paper says. "
     "Your job is to check whether the tool's findings are correct, and whether they matter. "
     "You do not review the paper itself and our tool does not claim to do that.", False),
    ("", False),
    ("We audited papers of NeurIPS 2025. The NeurIPS Paper Checklist Guidelines "
     "(https://neurips.cc/public/guides/PaperChecklist) specify that main experimental results "
     "include the new method and baselines, that authors should try to capture as many of the "
     "minor experiments in the paper as possible, and that if only a subset of experiments are "
     "reproducible, the paper should state which ones are. They also note that the instructions "
     "should contain the exact command and environment needed to reproduce the results.", False),
    ("", False),
    ("For each paper assigned to you there is one sheet in this workbook (tab = paper number). "
     "Alongside this workbook you received, per paper, the PDF and a copy of the released code. "
     "Each row of a paper sheet is one issue: the tool's claim, why it thinks it matters, the "
     "file and line, and the quoted code.", False),
    ("", False),
    ("HOW TO WORK", True),
    ("1. Enter your name at the top of each of your paper sheets.", False),
    ("2. Work through the issues one by one. For each issue, please note the minutes you worked "
     "on it until you came to a conclusion (rounding to minutes is fine), including reading the "
     "finding and filling the row.", False),
    ("3. Fill the two verdict columns (dropdowns) and, where required, the reason field.", False),
    ("", False),
    ("AXIS 1, CORRECTNESS: open the cited file and line and check the claim yourself, and "
     "compare to what the paper says if necessary.", True),
    ("True: the discrepancy exists as described. The claim is factually true (even if it is a "
     "nitpick).", False),
    ("Partially true: something real is there, but some part of it is wrong.", False),
    ("False: the claim doesn't hold, the tool misread the code or wrongly flags clearly "
     "intended, justified behavior.", False),
    ("Unverifiable: you can't tell (e.g. needs compute, you are missing expertise, you do not "
     "understand this part of the code).", False),
    ("", False),
    ("AXIS 2, RELEVANCE: do you think it would benefit reproducibility if this were fixed? Is "
     "there still a reasonable way for a reader to reproduce or verify the affected result "
     "without too much nuisance and work? Is the code quality, in your opinion, acceptable for "
     "a NeurIPS paper (high impact research code)?", True),
    ("Relevant: influences reproducibility or correctness of a main result.", False),
    ("Minor: worth a minor point in a review, but results don't depend on it or a reader can "
     "work around it with low effort.", False),
    ("Trivial: correct but easily correctable (e.g. no script re-plots a figure, but the "
     "numbers behind it are shipped in a results file).", False),
    ("Unverifiable: you can't tell.", False),
    ("N/A: use when Correctness is False.", False),
    ("", False),
    ("RULES", True),
    ("Please write a short reason whenever you answer False, Trivial, or Unverifiable.", False),
    ("Only use the provided code: the authors might have new code online in the meantime, but "
     "we did not audit this with the tool.", False),
    ("Please work alone and do not discuss issues with the other reviewers until everything is "
     "collected (some papers are deliberately reviewed by two people and we measure agreement).", False),
    ("Questions about the procedure: ask the study coordinator, not the other reviewers.", False),
]

HDR = ["#", "Issue id", "Category", "Location", "Paper reference", "Claim",
       "Why the tool thinks it matters", "Quoted code (from the cited location)",
       "Minutes worked", "Correctness",
       "Reason correctness (required for False / Unverifiable)", "Relevance",
       "Reason relevance (required for Trivial / Unverifiable)", "Notes (optional)"]
WIDTHS = [4, 22, 12, 30, 26, 60, 60, 50, 10, 16, 40, 22, 40, 28]

HEAD_FILL = PatternFill("solid", fgColor="DDE6F2")
INPUT_FILL = PatternFill("solid", fgColor="FFF7DD")
WRAP = Alignment(wrap_text=True, vertical="top")


def load_selection():
    a = json.load(open(os.path.join(SEL, "assignment.json")))
    covered = sorted(a["covered"], key=int)
    papers = {}
    for f in glob.glob(os.path.join(ROOT, "audits/*/findings_verified.json")):
        pdir = os.path.basename(os.path.dirname(f))
        pid = pdir.split("_")[0]
        if pid not in covered:
            continue
        title = ""
        for line in open(os.path.join(ROOT, "audits", pdir, "metadata.txt")):
            if line.startswith("title:"):
                title = line.split(":", 1)[1].strip()
                break
        items = [it for it in json.load(open(f))["findings"]
                 if it.get("verdict") in ("keep", "lowered")
                 and (it["severity"], it["category"]) in IN_SCOPE]
        items.sort(key=lambda it: (it.get("file") or "", it.get("line_start") or 0, it["id"]))
        papers[pid] = (title, items)
    return [(pid, *papers[pid]) for pid in covered]


def build():
    wb = Workbook()
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 110
    for i, (text, bold) in enumerate(INSTRUCTIONS, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.alignment = WRAP
        if bold:
            c.font = Font(bold=True)

    # assignment / progress table
    a = json.load(open(os.path.join(SEL, "assignment.json")))
    owner = {p: h for h, ps in a["heavy"].items() for p in ps}
    second = {p: l for l, ps in a["light"].items() for p in ps}
    r0 = len(INSTRUCTIONS) + 2
    ws.cell(row=r0, column=1, value="WHO DOES WHAT (fill in real names; every sheet below "
            "belongs to one paper)").font = Font(bold=True)
    for j, (h, w) in enumerate(zip(["Paper", "Title", "Issues", "Reviewer 1", "Reviewer 2"],
                                   [8, 70, 8, 18, 18]), 2):
        c = ws.cell(row=r0 + 1, column=j, value=h)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        ws.column_dimensions[get_column_letter(j)].width = w
    for i, (pid, title, items) in enumerate(load_selection()):
        r = r0 + 2 + i
        ws.cell(row=r, column=2, value=int(pid))
        ws.cell(row=r, column=3, value=title).alignment = WRAP
        ws.cell(row=r, column=4, value=len(items))
        ws.cell(row=r, column=5, value=owner.get(pid, "")).fill = INPUT_FILL
        ws.cell(row=r, column=6, value=second.get(pid, "")).fill = INPUT_FILL

    for pid, title, items in load_selection():
        ws = wb.create_sheet(str(pid))
        ws.cell(row=1, column=1, value=f"Paper #{pid}: {title}").font = Font(bold=True, size=12)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
        ws.cell(row=2, column=1, value="Reviewer name:").font = Font(bold=True)
        ws.cell(row=2, column=2).fill = INPUT_FILL
        for j, (h, w) in enumerate(zip(HDR, WIDTHS), 1):
            c = ws.cell(row=4, column=j, value=h)
            c.font = Font(bold=True)
            c.fill = HEAD_FILL
            c.alignment = WRAP
            ws.column_dimensions[get_column_letter(j)].width = w
        dv_c = DataValidation(type="list", formula1='"' + ",".join(CORRECTNESS) + '"',
                              allow_blank=True, showDropDown=False)
        dv_r = DataValidation(type="list", formula1='"' + ",".join(RELEVANCE) + '"',
                              allow_blank=True, showDropDown=False)
        ws.add_data_validation(dv_c)
        ws.add_data_validation(dv_r)
        for n, it in enumerate(items, 1):
            r = 4 + n
            lines = f"{it.get('line_start', '?')}"
            if it.get("line_end") and it["line_end"] != it.get("line_start"):
                lines += f"-{it['line_end']}"
            loc = f"{it.get('file', '?')}:{lines}"
            quote = loc + "\n" + (it.get("quote") or "").rstrip("\n")
            vals = [n, it.get("id_local") or it["id"].split("/")[-1],
                    CATEGORY_LABEL.get(it["category"], it["category"]),
                    loc,
                    it.get("paper_ref") or "", it.get("claim", "").strip(),
                    it.get("concern", "").strip(), quote,
                    None, None, None, None, None, None]
            for j, v in enumerate(vals, 1):
                c = ws.cell(row=r, column=j, value=v)
                c.alignment = WRAP
            for j in (9, 10, 11, 12, 13, 14):
                ws.cell(row=r, column=j).fill = INPUT_FILL
            dv_c.add(ws.cell(row=r, column=10))
            dv_r.add(ws.cell(row=r, column=12))
            ws.row_dimensions[r].height = 130
        ws.freeze_panes = "A5"
    wb.save(OUT)
    print("wrote", OUT, f"({len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    build()
